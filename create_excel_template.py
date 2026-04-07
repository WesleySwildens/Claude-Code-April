import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Financial Overview"

# Config
markets = ["SME", "Mid Market"]
editions = ["Premium", "Professional", "Enterprise"]
service_levels = ["SaaS", "MPS", "BPO"]
years = [2026, 2027, 2028, 2029, 2030, 2031]
metrics = ["Amount", "Price", "Total"]

# Styles
header_font = Font(bold=True, size=12)
section_font = Font(bold=True, size=13, color="FFFFFF")
subsection_font = Font(bold=True, size=11)
metric_font = Font(size=10)
total_font = Font(bold=True, size=10)
currency_format = '#,##0.00'
number_format = '#,##0'

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

market_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
edition_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
total_row_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
subtotal_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
grand_total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

# Column layout: A=Edition, B=Service Level, C=Metric, D-I=Years (2026-2031)
year_start_col = 4  # column D

# Set column widths
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 10
for i, yr in enumerate(years):
    ws.column_dimensions[get_column_letter(year_start_col + i)].width = 15

# Header row - Years
row = 1
ws.cell(row=row, column=1, value="Software Edition").font = header_font
ws.cell(row=row, column=2, value="Service Level").font = header_font
ws.cell(row=row, column=3, value="Metric").font = header_font
for i, yr in enumerate(years):
    cell = ws.cell(row=row, column=year_start_col + i, value=yr)
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

# Apply border to header
for col in range(1, year_start_col + len(years)):
    ws.cell(row=row, column=col).border = thin_border

current_row = 2

# Track total rows per market for grand total
market_total_rows = {}

for market in markets:
    # Market header row
    for col in range(1, year_start_col + len(years)):
        cell = ws.cell(row=current_row, column=col)
        cell.fill = market_fill
        cell.font = section_font
        cell.border = thin_border
    ws.cell(row=current_row, column=1, value=market)
    current_row += 1

    combo_total_rows = []  # track "Total" rows for subtotal

    for edition in editions:
        first_edition_row = current_row
        for sl in service_levels:
            for m_idx, metric in enumerate(metrics):
                cell_a = ws.cell(row=current_row, column=1)
                cell_b = ws.cell(row=current_row, column=2)
                cell_c = ws.cell(row=current_row, column=3, value=metric)

                # Only show edition/service level label on first metric row
                if m_idx == 0:
                    cell_a.value = edition
                    cell_a.font = subsection_font
                    cell_b.value = sl
                    cell_b.font = subsection_font

                cell_a.border = thin_border
                cell_b.border = thin_border
                cell_c.border = thin_border
                cell_c.font = metric_font

                for y_idx in range(len(years)):
                    col = year_start_col + y_idx
                    cell = ws.cell(row=current_row, column=col)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='right')

                    if metric == "Amount":
                        cell.number_format = number_format
                    elif metric == "Price":
                        cell.number_format = currency_format
                    elif metric == "Total":
                        # Formula: Amount * Price (2 rows up * 1 row up)
                        amount_cell = f"{get_column_letter(col)}{current_row - 2}"
                        price_cell = f"{get_column_letter(col)}{current_row - 1}"
                        cell.value = f"={amount_cell}*{price_cell}"
                        cell.number_format = currency_format
                        cell.font = total_font
                        cell.fill = total_row_fill

                if metric == "Total":
                    combo_total_rows.append(current_row)

                current_row += 1

    # Subtotal row for this market
    subtotal_row = current_row
    ws.cell(row=current_row, column=1, value=f"{market} Subtotal").font = Font(bold=True, size=11)
    for col_idx in range(1, year_start_col + len(years)):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.fill = subtotal_fill
        cell.border = thin_border
        cell.font = Font(bold=True, size=11)

    for y_idx in range(len(years)):
        col = year_start_col + y_idx
        cell = ws.cell(row=current_row, column=col)
        # SUM of all Total rows for this market
        parts = [f"{get_column_letter(col)}{r}" for r in combo_total_rows]
        cell.value = f"={'+'.join(parts)}"
        cell.number_format = currency_format
        cell.alignment = Alignment(horizontal='right')

    market_total_rows[market] = subtotal_row
    current_row += 2  # blank row between markets

# Grand Total row
ws.cell(row=current_row, column=1, value="Grand Total").font = Font(bold=True, size=12)
for col_idx in range(1, year_start_col + len(years)):
    cell = ws.cell(row=current_row, column=col_idx)
    cell.fill = grand_total_fill
    cell.border = thin_border
    cell.font = Font(bold=True, size=12)

for y_idx in range(len(years)):
    col = year_start_col + y_idx
    cell = ws.cell(row=current_row, column=col)
    col_letter = get_column_letter(col)
    parts = [f"{col_letter}{r}" for r in market_total_rows.values()]
    cell.value = f"={'+'.join(parts)}"
    cell.number_format = currency_format
    cell.alignment = Alignment(horizontal='right')

# Freeze panes (freeze header row + first 3 columns)
ws.freeze_panes = "D2"

output_path = "/home/user/Claude-Code-April/Financial_Template.xlsx"
wb.save(output_path)
print(f"Saved to {output_path}")
