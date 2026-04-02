#!/usr/bin/env python3
"""Generate KPI-Driven Operations Pitch Deck. Clean modern design, no template."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = "/home/user/Claude-Code-April/KPI_Driven_Operations_Pitch_Deck.pptx"

# Colors - SD Worx inspired but softer
NAVY = RGBColor(0x2B, 0x30, 0x3A)
BLUE = RGBColor(0x3D, 0x85, 0xB0)
SLATE = RGBColor(0x6E, 0x7B, 0x8F)
LIGHT = RGBColor(0xF2, 0xF3, 0xF5)
ORANGE = RGBColor(0xF0, 0x55, 0x1E)
PURPLE = RGBColor(0x6E, 0x0A, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x1A, 0x8C, 0x8C)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]  # blank layout


def txt(slide, l, t, w, h, text, sz=14, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    box.text_frame.word_wrap = True
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Roboto"
    return box


def multi(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT):
    """lines = [(text, size, bold, color), ...]"""
    box = slide.shapes.add_textbox(l, t, w, h)
    box.text_frame.word_wrap = True
    for i, (text, sz, bold, color) in enumerate(lines):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Roboto"


def card(slide, l, t, w, h, title, bullets, bg=NAVY, tc=WHITE, bc=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = bg
    s.line.fill.background()
    s.adjustments[0] = 0.06
    txt(slide, l + Inches(0.35), t + Inches(0.3), w - Inches(0.7), Inches(0.5),
        title, sz=18, bold=True, color=tc)
    lines = [(b, 13, False, bc) for b in bullets]
    multi(slide, l + Inches(0.35), t + Inches(0.9), w - Inches(0.7), h - Inches(1.1), lines)


def title_sub(slide, title, sub):
    txt(slide, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8),
        title, sz=28, bold=True, color=NAVY)
    txt(slide, Inches(0.9), Inches(1.2), Inches(11), Inches(0.5),
        sub, sz=14, color=SLATE)


# ── SLIDE 1: Titel ──
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.4),
                         Inches(12.5), Inches(6.7))
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.adjustments[0] = 0.03
txt(s, Inches(1.2), Inches(2.0), Inches(10), Inches(1.5),
    "Driving Value Through\nKPI-Driven Operations", sz=42, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
# orange accent line
ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(3.9), Inches(2.5), Pt(4))
ln.fill.solid()
ln.fill.fore_color.rgb = ORANGE
ln.line.fill.background()
txt(s, Inches(1.2), Inches(4.2), Inches(10), Inches(0.6),
    "Van inzicht naar actie per afdeling", sz=20, color=RGBColor(0xBB, 0xBF, 0xC7))
txt(s, Inches(1.2), Inches(6.0), Inches(3), Inches(0.4), "2026", sz=14, color=SLATE)

# ── SLIDE 2: Kernboodschap ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Waarom KPI-gedreven werken?", "De kernboodschap")
card(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.5),
     "", [
         "Zonder duidelijke KPI's  -->  geen focus",
         "",
         "Zonder focus  -->  geen voorspelbare resultaten",
         "",
         "Zonder resultaten  -->  geen schaalbare groei",
     ], bg=NAVY, tc=WHITE, bc=RGBColor(0xCC, 0xCF, 0xD5))
txt(s, Inches(1.3), Inches(5.2), Inches(10), Inches(0.5),
    "Elk team moet sturen op meetbare waarde", sz=20, bold=True, color=ORANGE)

# ── SLIDE 3: Ons uitgangspunt ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Een manier van werken, meerdere teams", "Ons uitgangspunt")
card(s, Inches(0.9), Inches(2.0), Inches(5.5), Inches(4.5),
     "Elk team heeft:", [
         "Een duidelijke doelstelling",
         "Een set KPI's",
         "Inzicht in bijdrage aan totaalresultaat",
     ], bg=BLUE)
card(s, Inches(6.9), Inches(2.0), Inches(5.5), Inches(4.5),
     "KPI's zijn:", [
         "Beinvloedbaar door het team",
         "Meetbaar en concreet",
         "Gekoppeld aan waarde",
     ], bg=PURPLE)

# ── SLIDE 4: Top-down + Bottom-up ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Balans tussen centrale sturing en team ownership", "Top-down + Bottom-up")
card(s, Inches(0.9), Inches(2.0), Inches(5.5), Inches(4.0),
     "Top-down KPI's", [
         "Strategische doelen organisatie",
         "Financiele performance",
         "Klantimpact",
     ], bg=NAVY)
card(s, Inches(6.9), Inches(2.0), Inches(5.5), Inches(4.0),
     "Bottom-up KPI's", [
         "Team-specifieke metrics",
         "Operationele drivers",
         "Dagelijkse sturing",
     ], bg=TEAL)
# bottom callout
cb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(6.3),
                         Inches(6.3), Inches(0.7))
cb.fill.solid()
cb.fill.fore_color.rgb = LIGHT
cb.line.fill.background()
cb.adjustments[0] = 0.3
txt(s, Inches(3.7), Inches(6.38), Inches(5.9), Inches(0.5),
    "Samen vormen ze een KPI-structuur", sz=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# ── SLIDE 5: Van strategie naar operatie (flow) ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Van strategie naar operatie", "Elke KPI moet leiden tot concreet gedrag")
flow = [
    ("Strategie", "Organisatie-\ndoelstellingen", PURPLE),
    ("Centrale KPI's", "Financieel,\nklant, groei", NAVY),
    ("Team KPI's", "Per afdeling\nmeetbaar", BLUE),
    ("Dagelijkse Acties", "Concreet\ngedrag", ORANGE),
]
bw = Inches(2.7)
bh = Inches(3.5)
gap = Inches(0.45)
aw = Inches(0.3)
sx = Inches(0.6)
for i, (t, d, c) in enumerate(flow):
    x = sx + i * (bw + gap + aw)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), bw, bh)
    sh.fill.solid()
    sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    sh.adjustments[0] = 0.06
    txt(s, x + Inches(0.2), Inches(2.6), bw - Inches(0.4), Inches(0.6),
        t, sz=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # divider
    dv = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.6), Inches(3.4), bw - Inches(1.2), Pt(2))
    dv.fill.solid()
    dv.fill.fore_color.rgb = WHITE
    dv.line.fill.background()
    txt(s, x + Inches(0.2), Inches(3.7), bw - Inches(0.4), Inches(1.2),
        d, sz=15, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                x + bw + Inches(0.07), Inches(3.7), aw, Inches(0.35))
        ar.fill.solid()
        ar.fill.fore_color.rgb = SLATE
        ar.line.fill.background()

# ── SLIDE 6: Recurring Services ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Recurring Services", "Sturen op klantwaarde en efficiency")
cw = Inches(3.7)
ch = Inches(4.0)
cg = Inches(0.45)
rs_data = [
    ("NPS", ["Klanttevredenheid", "Klantbehoud versterken", "Signalen vroeg oppakken"], BLUE),
    ("eNPS", ["Medewerkerbetrokkenheid", "Teamvitaliteit meten", "Retentie waarborgen"], NAVY),
    ("Contribution /\nCost to Serve", ["Efficiente service delivery", "Schaalbaarheid", "Waarde per klant optimaliseren"], PURPLE),
]
for i, (t, b, c) in enumerate(rs_data):
    card(s, Inches(0.9) + i * (cw + cg), Inches(2.0), cw, ch, t, b, bg=c)

# ── SLIDE 7: Consulting ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Consulting", "Sturen op waardecreatie")
con = [
    ("tNPS", ["Onboarding / implementatie", "Klanttevredenheid bij trajectory", "Kwaliteit van oplevering"], BLUE),
    ("eNPS", ["Medewerkerbetrokkenheid", "Teamdynamiek en groei", "Consultants als ambassadeurs"], NAVY),
    ("Revenue per FTE /\nBillability", ["Productiviteit per consultant", "Effectieve inzet capaciteit", "Financiele gezondheid team"], TEAL),
    ("Proven Value Time", ["(Toekomst) Tijd met/voor klant", "Aantoonbare waardecreatie", "Meetbare klantimpact"], ORANGE),
]
cw2 = Inches(2.85)
for i, (t, b, c) in enumerate(con):
    card(s, Inches(0.7) + i * (cw2 + Inches(0.3)), Inches(2.0), cw2, Inches(4.0), t, b, bg=c)
# bottom callout
cb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(6.3),
                         Inches(8.3), Inches(0.7))
cb.fill.solid()
cb.fill.fore_color.rgb = LIGHT
cb.line.fill.background()
cb.adjustments[0] = 0.3
txt(s, Inches(2.7), Inches(6.38), Inches(7.9), Inches(0.5),
    "Belangrijke shift: van uren schrijven naar aantoonbare klantwaarde",
    sz=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# ── SLIDE 8: KC ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "KC - Knowledge Center", "Sturen op kwaliteit en adoptie")
kc = [
    ("Kennis & Kwaliteit", ["Kennisborging", "Kwaliteitsstandaarden", "Kennisdeling bevorderen"], BLUE),
    ("Compliance & Control", ["Risicobeheersing", "Naleving van regelgeving", "Audit-readiness"], NAVY),
    ("Training & Adoption", ["Gebruik van oplossingen", "Adoptiegraad verhogen", "Continue ontwikkeling"], PURPLE),
]
for i, (t, b, c) in enumerate(kc):
    card(s, Inches(0.9) + i * (cw + cg), Inches(2.0), cw, ch, t, b, bg=c)

# ── SLIDE 9: Ownership ──
s = prs.slides.add_slide(BLANK)
title_sub(s, "Ownership bij de teams", "Wat verwachten we van elk team?")
own = [
    ("Begrijp je KPI's", "Weet wat je meet en\nwaarom het belangrijk is", PURPLE),
    ("Weet hoe je ze\nbeinvloedt", "Ken de hefbomen die\nje als team kunt gebruiken", NAVY),
    ("Stuur actief bij", "Wacht niet af maar neem\ninitiatief om bij te sturen", BLUE),
    ("Maak impact\nzichtbaar", "Laat resultaten zien\nen deel successen", ORANGE),
]
pw = Inches(2.8)
ph = Inches(4.2)
for i, (t, d, c) in enumerate(own):
    x = Inches(0.7) + i * (pw + Inches(0.35))
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), pw, ph)
    sh.fill.solid()
    sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    sh.adjustments[0] = 0.06
    # number circle
    ci = s.shapes.add_shape(MSO_SHAPE.OVAL, x + pw/2 - Inches(0.35), Inches(2.4),
                            Inches(0.7), Inches(0.7))
    ci.fill.solid()
    ci.fill.fore_color.rgb = WHITE
    ci.line.fill.background()
    txt(s, x + pw/2 - Inches(0.35), Inches(2.47), Inches(0.7), Inches(0.6),
        f"0{i+1}", sz=20, bold=True, color=c, align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.2), Inches(3.4), pw - Inches(0.4), Inches(0.8),
        t, sz=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.2), Inches(4.3), pw - Inches(0.4), Inches(1.2),
        d, sz=14, color=WHITE, align=PP_ALIGN.CENTER)

# ── Save ──
prs.save(OUTPUT)
print(f"Saved: {OUTPUT} ({len(prs.slides)} slides)")
