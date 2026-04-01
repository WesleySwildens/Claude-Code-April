#!/usr/bin/env python3
"""
Generate a Vertical SaaS Guide PowerPoint using the SD Worx template.
Uses the template's slide layouts, theme colors, and fonts.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy

TEMPLATE = "/home/user/Claude-Code-April/2602 - SD Worx PowerPoint GOED .pptx"
OUTPUT = "/home/user/Claude-Code-April/Vertical_Approach_SaaS_Guide.pptx"

# SD Worx theme colors
DK1 = RGBColor(0x30, 0x36, 0x42)      # Dark navy
ACCENT2 = RGBColor(0x43, 0x8A, 0xB5)  # Blue
ACCENT3 = RGBColor(0x75, 0x82, 0x9B)  # Slate gray
ACCENT4 = RGBColor(0xE4, 0xE6, 0xEC)  # Light gray
ACCENT5 = RGBColor(0xFF, 0x4E, 0x0F)  # SD Worx Orange
ACCENT6 = RGBColor(0x7A, 0x00, 0x51)  # SD Worx Purple/Magenta
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

# Load template
prs = Presentation(TEMPLATE)

# Map layout names to layout objects
layout_map = {}
for layout in prs.slide_layouts:
    layout_map[layout.name] = layout

# Remove all existing slides
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

# ── Helper functions ──
def get_layout(name):
    """Get a layout by name, with fallback."""
    if name in layout_map:
        return layout_map[name]
    # Fallback search
    for k, v in layout_map.items():
        if name.lower() in k.lower():
            return v
    return prs.slide_layouts[0]


def set_placeholder_text(slide, idx, text, font_size=None, bold=None, color=None):
    """Set text on a placeholder by index, if it exists."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            if ph.text_frame.paragraphs:
                para = ph.text_frame.paragraphs[0]
                if font_size or bold is not None or color:
                    for run in para.runs:
                        if font_size:
                            run.font.size = Pt(font_size)
                        if bold is not None:
                            run.font.bold = bold
                        if color:
                            run.font.color.rgb = color
            return ph
    return None


def set_placeholder_bullets(slide, idx, items, font_size=14, color=None):
    """Set bulleted text on a placeholder."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            tf = ph.text_frame
            tf.clear()
            for i, item in enumerate(items):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = item
                p.font.size = Pt(font_size)
                if color:
                    p.font.color.rgb = color
                p.space_after = Pt(6)
            return ph
    return None


def add_content_slide(layout_name, title, subtitle, body_items, font_size=14):
    """Add a standard content slide."""
    slide = prs.slides.add_slide(get_layout(layout_name))
    set_placeholder_text(slide, 0, title)
    set_placeholder_text(slide, 14, subtitle)
    set_placeholder_bullets(slide, 1, body_items, font_size=font_size)
    return slide


# ═══════════════════════════════════════════════════════════════════
# SLIDE 1 – Cover / Title
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Cover text - option 1'))
set_placeholder_text(slide, 0, "The Vertical Approach\nin SaaS")
set_placeholder_text(slide, 10, "A Strategic Guide to Building Industry-Specific Software Solutions")
set_placeholder_text(slide, 12, "What every founder, product leader, and GTM team needs to know")
set_placeholder_text(slide, 13, "2026")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 2 – Agenda / Overview
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Agenda - option 1'))
set_placeholder_text(slide, 0, "Agenda")
set_placeholder_text(slide, 14, "What we will cover today")
set_placeholder_bullets(slide, 16, [
    "1\tWhat is Vertical SaaS?",
    "2\tWhy Go Vertical? Strategic Advantages",
    "3\tChoosing the Right Vertical",
    "4\tProduct Strategy: Building for the Vertical",
    "5\tGo-to-Market Playbook",
    "6\tKey Metrics & Benchmarks",
    "7\tCommon Pitfalls to Avoid",
    "8\tScaling & Expansion Framework",
    "9\tKey Takeaways",
], font_size=16)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 3 – Chapter: What is Vertical SaaS?
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 1'))
set_placeholder_text(slide, 11, "01")
set_placeholder_text(slide, 0, "What Is a Vertical\nSaaS Approach?")
set_placeholder_text(slide, 10, "Understanding the model and how it differs from horizontal SaaS")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 4 – What is Vertical SaaS (2-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 2 blocks'))
set_placeholder_text(slide, 0, "Vertical vs. Horizontal SaaS")
set_placeholder_text(slide, 18, "Two fundamentally different approaches to building software")

# Left block
set_placeholder_text(slide, 19, "Vertical SaaS", bold=True)
set_placeholder_bullets(slide, 1, [
    "Software built for ONE specific industry or niche",
    "Deeply tailored workflows, compliance & terminology",
    "End-to-end solution replacing horizontal point tools",
    "Examples: Veeva (pharma), Procore (construction), Toast (restaurants)",
    "Smaller TAM but higher win rate and stickiness",
], font_size=13)

# Right block
set_placeholder_text(slide, 22, "Horizontal SaaS", bold=True)
set_placeholder_bullets(slide, 23, [
    "Software built for any industry or use case",
    "Broad feature set, generic terminology and UX",
    "Point solution for a specific function (CRM, HR, etc.)",
    "Examples: Salesforce, HubSpot, Slack, Asana",
    "Larger TAM but more competition and lower retention",
], font_size=13)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 5 – Chapter: Why Go Vertical?
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 2'))
set_placeholder_text(slide, 11, "02")
set_placeholder_text(slide, 0, "Why Go Vertical?")
set_placeholder_text(slide, 10, "The strategic advantages of a vertical SaaS approach")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 6 – Why Go Vertical (4-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 4 blocks option 1'))
set_placeholder_text(slide, 0, "Strategic Advantages of Vertical SaaS")
set_placeholder_text(slide, 18, "Four key reasons why going vertical wins")

# Block titles (with number indicators)
set_placeholder_text(slide, 33, "1")
set_placeholder_text(slide, 19, "Higher Win Rates", bold=True)
set_placeholder_bullets(slide, 29, [
    "Prospects see you as the expert",
    "You speak their language and understand their pain",
    "Demo workflows they recognize immediately",
], font_size=11)

set_placeholder_text(slide, 34, "2")
set_placeholder_text(slide, 28, "Stronger Retention", bold=True)
set_placeholder_bullets(slide, 30, [
    "Deep integration = high switching costs",
    "Vertical SaaS often sees 120-140%+ NRR",
    "Becomes system of record for the business",
], font_size=11)

set_placeholder_text(slide, 35, "3")
set_placeholder_text(slide, 25, "Efficient GTM", bold=True)
set_placeholder_bullets(slide, 17, [
    "Concentrated buyer personas",
    "Focused conferences & trade publications",
    "Word-of-mouth referrals spread fast",
], font_size=11)

set_placeholder_text(slide, 36, "4")
set_placeholder_text(slide, 32, "Pricing Power", bold=True)
set_placeholder_bullets(slide, 31, [
    "Industry-specific value justifies premium pricing",
    "Customers pay for outcomes, not features",
    "Ability to layer on fintech, data, services",
], font_size=11)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 7 – Chapter: Choosing Your Vertical
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 1'))
set_placeholder_text(slide, 11, "03")
set_placeholder_text(slide, 0, "Choosing the\nRight Vertical")
set_placeholder_text(slide, 10, "Selection criteria for identifying the best industry to target")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 8 – Selection Criteria (content)
# ═══════════════════════════════════════════════════════════════════
slide = add_content_slide(
    'Content',
    "Vertical Selection Criteria",
    "Seven key factors to evaluate before committing to a vertical",
    [
        "Market Size & Density — Is the TAM large enough? Are buyers concentrated or fragmented?",
        "Underserved by Tech — Are incumbents still using spreadsheets, paper, or legacy on-prem systems?",
        "Regulatory Complexity — Industries with compliance needs (HIPAA, SOX, FDA) create natural moats",
        "Willingness to Pay — Does the industry have healthy margins? Do they already budget for software?",
        "Domain Access — Do you have founder-market fit? Can you access early design partners?",
        "Workflow Standardization — Are core workflows consistent enough to build a scalable product?",
        "Expansion Potential — Can you layer on payments, financing, marketplace, or data products?",
    ],
    font_size=14
)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 9 – Chapter: Product Strategy
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 2'))
set_placeholder_text(slide, 11, "04")
set_placeholder_text(slide, 0, "Product Strategy")
set_placeholder_text(slide, 10, "Building for the vertical: domain modeling, workflows, and platform")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 10 – Product Strategy (3-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 3 blocks'))
set_placeholder_text(slide, 0, "Three Pillars of Vertical Product Strategy")
set_placeholder_text(slide, 18, "How to build a product that dominates its vertical")

set_placeholder_text(slide, 19, "Deep Domain Modeling", bold=True)
set_placeholder_bullets(slide, 1, [
    "Map the industry value chain end to end",
    "Build domain-specific data models & objects",
    "Use industry terminology in UX",
    "Embed compliance rules into the product",
], font_size=12)

set_placeholder_text(slide, 28, "Workflow-First Design", bold=True)
set_placeholder_bullets(slide, 29, [
    "Shadow real users in their environment",
    "Digitize existing paper/manual processes",
    "Automate boring, error-prone steps",
    "Design for the least technical user",
], font_size=12)

set_placeholder_text(slide, 25, "Platform & Ecosystem", bold=True)
set_placeholder_bullets(slide, 30, [
    "Integrate with industry-specific ERPs",
    "Build APIs for partner integrations",
    "Create analytics & benchmarking layer",
    "Plan for embedded fintech",
], font_size=12)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 11 – Chapter: Go-to-Market
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 1'))
set_placeholder_text(slide, 11, "05")
set_placeholder_text(slide, 0, "Go-to-Market\nPlaybook")
set_placeholder_text(slide, 10, "Selling, expanding, and growing in a focused vertical")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 12 – GTM Playbook (4-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 4 blocks option 1'))
set_placeholder_text(slide, 0, "The Vertical SaaS GTM Playbook")
set_placeholder_text(slide, 18, "Four phases of a winning go-to-market motion")

set_placeholder_text(slide, 33, "1")
set_placeholder_text(slide, 19, "Land with a Wedge", bold=True)
set_placeholder_bullets(slide, 29, [
    "Start with ONE critical pain point",
    "Win a single workflow before expanding",
    "Early adopters become evangelists",
], font_size=11)

set_placeholder_text(slide, 34, "2")
set_placeholder_text(slide, 28, "Industry-Native Sales", bold=True)
set_placeholder_bullets(slide, 30, [
    "Hire reps from the industry",
    "Attend vertical trade shows",
    "Publish thought leadership in trade publications",
], font_size=11)

set_placeholder_text(slide, 35, "3")
set_placeholder_text(slide, 25, "Customer-Led Growth", bold=True)
set_placeholder_bullets(slide, 17, [
    "Case studies & ROI calculators",
    "Peer referrals in tight communities",
    "Reputation spreads fast — good and bad",
], font_size=11)

set_placeholder_text(slide, 36, "4")
set_placeholder_text(slide, 32, "Expand Revenue", bold=True)
set_placeholder_bullets(slide, 31, [
    "Layer on modules, payments, data",
    "Professional services motion",
    "3-5x initial ACV over time",
], font_size=11)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 13 – Chapter: Key Metrics
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 2'))
set_placeholder_text(slide, 11, "06")
set_placeholder_text(slide, 0, "Key Metrics &\nBenchmarks")
set_placeholder_text(slide, 10, "What good looks like for vertical SaaS businesses")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 14 – Metrics (content)
# ═══════════════════════════════════════════════════════════════════
slide = add_content_slide(
    'Content',
    "Vertical SaaS Benchmarks",
    "Target metrics for a healthy vertical SaaS business",
    [
        "Net Revenue Retention: 120-140%+ — Best-in-class verticals expand aggressively within accounts",
        "Gross Margin: 70-80%+ — Can be lower if services-heavy; aim for software-like margins over time",
        "CAC Payback: 12-18 months — Efficient GTM thanks to a concentrated buyer base",
        "Logo Retention: 90-95%+ — High switching costs and deep workflow integration keep churn low",
        "Market Penetration: 10-30% — Realistic ceiling within a single vertical segment before expanding",
        "ACV Expansion: 3-5x — Layer on modules, payments, data, and services to grow deal sizes",
    ],
    font_size=14
)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 15 – Chapter: Common Pitfalls
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 1'))
set_placeholder_text(slide, 11, "07")
set_placeholder_text(slide, 0, "Common Pitfalls\nto Avoid")
set_placeholder_text(slide, 10, "Mistakes that derail vertical SaaS companies")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 16 – Pitfalls (2-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 2 blocks'))
set_placeholder_text(slide, 0, "Six Common Pitfalls in Vertical SaaS")
set_placeholder_text(slide, 18, "Learn from the mistakes of others")

set_placeholder_text(slide, 19, "Strategic Pitfalls", bold=True)
set_placeholder_bullets(slide, 1, [
    "Going Too Broad Too Soon — Resist adjacent verticals before dominating your first one. Depth beats breadth early on.",
    "Underestimating Compliance — Regulatory requirements aren't optional. Bake HIPAA, PCI, SOC 2, and industry regs in from day one.",
    "Neglecting Data & Analytics — Your aggregated data is a strategic asset. Benchmarking and insights can become a standalone revenue stream.",
], font_size=12)

set_placeholder_text(slide, 22, "Execution Pitfalls", bold=True)
set_placeholder_bullets(slide, 23, [
    "Ignoring Services Revenue — Many verticals expect onboarding, training, and customization. Build a services motion — don't treat it as a distraction.",
    "Building for Power Users Only — Your ICP may include non-technical users. If the UX requires training, you'll lose deals.",
    "Over-Customizing Per Client — Custom work for one customer erodes margins. Build configurable, not custom. Say no to one-off requests.",
], font_size=12)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 17 – Chapter: Scaling Framework
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 2'))
set_placeholder_text(slide, 11, "08")
set_placeholder_text(slide, 0, "Scaling & Expansion\nFramework")
set_placeholder_text(slide, 10, "From wedge product to industry platform")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 18 – Scaling (4-block)
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Content - 4 blocks option 1'))
set_placeholder_text(slide, 0, "The Four Phases of Vertical SaaS Growth")
set_placeholder_text(slide, 18, "A roadmap from first customer to ecosystem dominance")

set_placeholder_text(slide, 33, "1")
set_placeholder_text(slide, 19, "Phase 1: Wedge", bold=True)
set_placeholder_bullets(slide, 29, [
    "Single killer workflow",
    "5-20 design partners",
    "Prove ROI in one use case",
    "Achieve product-market fit",
], font_size=11)

set_placeholder_text(slide, 34, "2")
set_placeholder_text(slide, 28, "Phase 2: Suite", bold=True)
set_placeholder_bullets(slide, 30, [
    "Add adjacent workflows",
    "Become system of record",
    "Launch self-serve onboarding",
    "Hit $1-5M ARR",
], font_size=11)

set_placeholder_text(slide, 35, "3")
set_placeholder_text(slide, 25, "Phase 3: Platform", bold=True)
set_placeholder_bullets(slide, 17, [
    "APIs & integrations layer",
    "Embedded fintech",
    "Marketplace / app store",
    "Hit $10-30M ARR",
], font_size=11)

set_placeholder_text(slide, 36, "4")
set_placeholder_text(slide, 32, "Phase 4: Ecosystem", bold=True)
set_placeholder_bullets(slide, 31, [
    "Data & benchmarking products",
    "Adjacent vertical expansion",
    "M&A bolt-ons",
    "Scale to $50M+ ARR",
], font_size=11)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 19 – Chapter: Key Takeaways
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Chapter text - option 1'))
set_placeholder_text(slide, 11, "09")
set_placeholder_text(slide, 0, "Key Takeaways")
set_placeholder_text(slide, 10, "What to remember from this guide")


# ═══════════════════════════════════════════════════════════════════
# SLIDE 20 – Takeaways (content)
# ═══════════════════════════════════════════════════════════════════
slide = add_content_slide(
    'Content',
    "Key Takeaways",
    "Six principles for succeeding with a vertical SaaS approach",
    [
        "Go deep before going wide — dominate one vertical before expanding to others",
        "Build with domain experts, not just engineers — hire from the industry you serve",
        "Compliance and workflows ARE the product — they are not afterthoughts or features",
        "Your GTM is your moat — industry relationships and reputation compound over time",
        "Layer revenue streams — software, payments, data products, and professional services",
        "Vertical SaaS commands premium valuations (10-20x+ ARR) thanks to stickiness and expansion potential",
    ],
    font_size=15
)


# ═══════════════════════════════════════════════════════════════════
# SLIDE 21 – Closing / Quote
# ═══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(get_layout('Quote / statement - gradient option 1'))
set_placeholder_text(slide, 0, '"The riches are in the niches.\nVertical SaaS is the future of enterprise software."')
set_placeholder_text(slide, 10, "Thank you — Questions?")


# ── Save ──
prs.save(OUTPUT)
print(f"Presentation saved to: {OUTPUT}")
print(f"Total slides: {len(prs.slides)}")
